"""Reproduce the public synthetic Quay PPTX fixture; no product runtime dependency.
Producer: python-pptx 1.0.2, Pillow 12.3.0. ZIP metadata is fixed after authoring.
"""
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile, ZipInfo, ZIP_DEFLATED
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches

out = Path(__file__).parent
image = Image.new('RGB', (480, 240), 'white')
draw = ImageDraw.Draw(image)
draw.rectangle((20, 30, 450, 90), fill='#2f6478')
draw.rectangle((20, 140, 300, 200), fill='#82a5ad')
draw.text((30, 50), 'Pier A: 7 berths', fill='white')
draw.text((30, 160), 'Pier B: 4 berths', fill='black')
png = BytesIO(); image.save(png, format='PNG')
jpeg = BytesIO(); image.save(jpeg, format='JPEG', quality=85)
(out / 'quay-map.png').write_bytes(png.getvalue())
(out / 'quay-map.jpg').write_bytes(jpeg.getvalue())
prs = Presentation()
first = prs.slides.add_slide(prs.slide_layouts[6])
second = prs.slides.add_slide(prs.slide_layouts[6])
first._element.set('show', '0')
for slide in [first, second]:
    shape = slide.shapes.add_textbox(Inches(.5), Inches(.2), Inches(7), Inches(.5))
    shape.text = 'Dock review is pending.'
group = first.shapes.add_group_shape()
box = group.shapes.add_textbox(Inches(.5), Inches(1), Inches(6), Inches(1))
p = box.text_frame.paragraphs[0]
p.add_run().text = 'Inspect pier A'
p.add_line_break()
p.add_run().text = 'before cargo arrival.'
p2 = box.text_frame.add_paragraph()
p2.text = 'Display date: '
fld = OxmlElement('a:fld'); fld.set('id', '{00112233-4455-6677-8899-AABBCCDDEEFF}'); fld.set('type', 'datetime1')
t = OxmlElement('a:t'); t.text = '14 October'; fld.append(t); p2._p.append(fld)
box._element.nvSpPr.cNvPr.set('hidden', '1')
unknown = OxmlElement('a:unsupported'); p2._p.append(unknown)
table = first.shapes.add_table(2, 3, Inches(.5), Inches(2.3), Inches(7), Inches(1)).table
for row, values in enumerate([['Pier', 'Owner', 'State'], ['A', 'Leena', 'Pending']]):
    for col, value in enumerate(values): table.cell(row, col).text = value
for slide, x in [(first, .5), (second, .5)]:
    picture = slide.shapes.add_picture(BytesIO(png.getvalue()), Inches(x), Inches(4), width=Inches(3))
    picture._element.nvPicPr.cNvPr.set('title', 'Quay berth diagram')
    picture._element.nvPicPr.cNvPr.set('descr', 'Authored pier capacity illustration')
second.shapes.add_picture(BytesIO(jpeg.getvalue()), Inches(4), Inches(4), width=Inches(3))
first.notes_slide.notes_text_frame.text = 'Notes only: Leena must confirm the crane by 14 October.'
second.notes_slide.notes_text_frame.text = 'Do not treat a berth count as arrival authorization.'
other = OxmlElement('p:contentPart'); first._element.cSld.spTree.append(other)
raw = BytesIO(); prs.save(raw)
with ZipFile(BytesIO(raw.getvalue())) as source, ZipFile(out / 'quay.pptx', 'w') as target:
    for entry in source.infolist():
        name = entry.filename.replace('slide1.xml', 'slide9.xml')
        data = source.read(entry)
        if name.endswith(('.xml', '.rels')): data = data.replace(b'slide1.xml', b'slide9.xml')
        stable = ZipInfo(name, date_time=(2026, 10, 9, 0, 0, 0)); stable.compress_type = ZIP_DEFLATED
        target.writestr(stable, data)
