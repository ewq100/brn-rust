"""Benign producer-generated synthetic Office/MIME corpus; no private input."""
from pathlib import Path
from io import BytesIO
from email.message import EmailMessage
from email.policy import SMTP
import json, hashlib, zipfile
from PIL import Image, ImageDraw
from docx import Document
from docx.shared import Inches
from pptx import Presentation
from pptx.util import Inches as PInches
from pptx.enum.shapes import MSO_SHAPE
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from openpyxl import Workbook
root=Path(__file__).parent; f=root/'fixtures'; f.mkdir(exist_ok=True)
im=Image.new('RGB',(900,500),'white'); draw=ImageDraw.Draw(im)
draw.text((35,25),'Harbor pilot: measured water use (litres/day)',fill='black',font_size=30)
for y,label,value,col in [(110,'Before',120,'#B34D39'),(270,'After',72,'#1C8068')]:
 draw.text((35,y),label,fill='black',font_size=26); draw.rectangle((170,y,170+value*5,y+90),fill=col); draw.text((180+value*5,y+25),str(value),fill='black',font_size=30)
draw.text((35,420),'Result: 40% less water. Source: synthetic pilot, 6 Oct 2026.',fill='black',font_size=23)
im.save(f/'water-use.png')
d=Document(); d.add_heading('Harbor water-saving pilot',0)
d.add_paragraph('Decision: extend the pilot to Pier B on 15 October 2026. Owner: Mira. The chart is evidence, not decoration.')
d.sections[0].header.paragraphs[0].text='Internal synthetic briefing — review date 12 October 2026'
d.sections[0].footer.paragraphs[0].text='Footer constraint: budget capped at EUR 4,000.'
t=d.add_table(rows=1,cols=3); t.style='Light Shading Accent 1'
for c,s in zip(t.rows[0].cells,['Site','Before','After']): c.text=s
for s in [['Pier A','120 L/day','72 L/day'],['Pier B','Pending','Pending']]:
 for c,v in zip(t.add_row().cells,s): c.text=v
for name,title in [('primary-chart','Water use before and after'),('repeated-chart','Same chart for operations handoff')]:
 d.add_paragraph(title); sh=d.add_picture(str(f/'water-use.png'),width=Inches(5.7)); pr=sh._inline.docPr; pr.set('name',name); pr.set('title',title); pr.set('descr','120 to 72 litres/day, a 40% reduction')
d.save(f/'harbor.docx')
r=Presentation(); s=r.slides.add_slide(r.slide_layouts[6]);
box=s.shapes.add_textbox(PInches(.4),PInches(.2),PInches(9),PInches(.7)); box.text='Harbor decision and dependencies'
for i,title in enumerate(['Primary water chart','Repeated water chart']):
 sh=s.shapes.add_picture(str(f/'water-use.png'),PInches(.4+i*4.7),PInches(1.2),width=PInches(4.4)); pr=sh._element.nvPicPr.cNvPr; pr.set('title',title); pr.set('descr','Water savings evidence')
shape=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,PInches(.5),PInches(4),PInches(8.8),PInches(1.2)); shape.text='CRITICAL: installation may start only after valve inspection.'; shape._element.nvSpPr.cNvPr.set('title','Not a picture: dependency gate')
s.notes_slide.notes_text_frame.text='Presenter note: ask Mira to book the inspection by 12 October.'
s2=r.slides.add_slide(r.slide_layouts[6]); b=s2.shapes.add_textbox(PInches(.4),PInches(.2),PInches(9),PInches(.7)); b.text='Visual-only scheduling evidence: inspection capacity'
cd=CategoryChartData(); cd.categories=['12 October','13 October']; cd.add_series('Available inspection slots',[2,0]); s2.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,PInches(.6),PInches(1.2),PInches(8.5),PInches(5),cd)
r.save(f/'harbor.pptx')
w=Workbook(); w.active.append(['Out of scope forecast','EUR']); w.active.append(['Harbor budget',4000]); w.save(f/'forecast.xlsx')
for case,extra in [('single',False),('plural',True)]:
 m=EmailMessage(policy=SMTP); m['From']='Mira <mira@example.test>'; m['To']='Operations <ops@example.test>'; m['Subject']='Harbor pilot – review'; m['Date']='Tue, 06 Oct 2026 16:20:00 +0300'; m['Message-ID']=f'<harbor-{case}@example.test>'
 m.set_content('Please review the attached Harbor pilot. Extend to Pier B after inspection.'); m.add_alternative('<html><body><p>Water use fell 40%.</p><img src="cid:water-evidence"><p>See attached decision.</p></body></html>',subtype='html')
 m.get_payload()[1].add_related((f/'water-use.png').read_bytes(),maintype='image',subtype='png',cid='<water-evidence>',filename='inline-water.png')
 m.add_attachment((f/'harbor.docx').read_bytes(),maintype='application',subtype='vnd.openxmlformats-officedocument.wordprocessingml.document',filename='harbor.docx')
 if extra: m.add_attachment((f/'forecast.xlsx').read_bytes(),maintype='application',subtype='vnd.openxmlformats-officedocument.spreadsheetml.sheet',filename='forecast.xlsx')
 (f/f'{case}.eml').write_bytes(m.as_bytes())
# Four hostile packages exercise outer admission, independent of converter internals.
for kind in ['path','dtd','inflate','members']:
 with zipfile.ZipFile(f/f'hostile-{kind}.docx','w',zipfile.ZIP_DEFLATED) as z:
  if kind=='path': z.writestr('../escape',b'no')
  if kind=='dtd': z.writestr('word/document.xml',b'<!DOCTYPE x [<!ENTITY e SYSTEM "file:///etc/passwd">]><x>&e;</x>')
  if kind=='inflate': z.writestr('word/document.xml',b'A'*(17*1024*1024))
  if kind=='members':
   for i in range(513): z.writestr(f'parts/{i}',b'x')
manifest={p.name:{'bytes':p.stat().st_size,'sha256':hashlib.sha256(p.read_bytes()).hexdigest()} for p in sorted(f.iterdir())}
(root/'evidence'/'fixture-manifest.json').write_text(json.dumps({'producers':{'python-docx':'1.2.0','python-pptx':'1.0.2','Pillow':'12.3.0'},'files':manifest},indent=2)+'\n')
